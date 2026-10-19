"""
Reporte mensual de descuentos por nómina, por fondo.

Flujo:
1. Durante el mes el fondo aprueba (total o parcial) solicitudes.
2. Al cierre, el gestor revisa la vista previa del mes y descarga el archivo
   (CSV o Excel) que el fondo usa para aplicar el descuento en nómina.
3. "Cerrar corte" congela ese listado: las solicitudes quedan marcadas con el
   corte y no vuelven a salir en meses siguientes. Las aprobadas tarde (después
   de un cierre) salen en el siguiente corte, así ninguna se pierde.
"""

import csv
import io
from dataclasses import dataclass
from datetime import date, datetime, time
from decimal import Decimal

from django.conf import settings
from django.core.mail import EmailMessage
from django.db import IntegrityError, transaction
from django.db.models import Sum
from django.utils import timezone

from .models import CorteNomina, SolicitudDescuentoNomina

S = SolicitudDescuentoNomina

COLUMNAS = [
    ("periodo", "Periodo"),
    ("documento", "Documento"),
    ("nombre", "Nombre"),
    ("empresa", "Empresa"),
    ("finca", "Finca / punto de entrega"),
    ("numero_afiliado", "Número de afiliado"),
    ("pedido", "Pedido"),
    ("fecha_aprobacion", "Fecha de aprobación"),
    ("tipo", "Tipo de aprobación"),
    ("valor_pedido", "Valor del pedido"),
    ("valor_descontar", "Valor a descontar"),
    ("referencia_fondo", "Referencia del fondo"),
    ("fecha_entrega", "Fecha de entrega"),
    ("solicitud", "ID solicitud"),
]


class ReporteError(Exception):
    pass


@dataclass
class Resumen:
    cantidad: int
    total: Decimal


def periodo_desde_texto(texto):
    """'2026-10' -> date(2026, 10, 1)."""
    try:
        anio, mes = (int(p) for p in str(texto).split("-")[:2])
        return date(anio, mes, 1)
    except (TypeError, ValueError):
        raise ReporteError("Mes inválido. Use el formato AAAA-MM, por ejemplo 2026-10.")


def periodo_por_defecto():
    """El mes anterior: el corte se hace al terminar el mes."""
    hoy = timezone.localdate()
    return date(hoy.year - 1, 12, 1) if hoy.month == 1 else date(hoy.year, hoy.month - 1, 1)


def _siguiente_mes(periodo):
    return date(periodo.year + 1, 1, 1) if periodo.month == 12 else date(periodo.year, periodo.month + 1, 1)


def solicitudes_pendientes_de_corte(fondo, periodo):
    """Aprobadas hasta el fin del mes, sin corte y con el pedido vigente."""
    from apps.pedidos.models import Pedido

    fin = timezone.make_aware(datetime.combine(_siguiente_mes(periodo), time.min))
    return (
        S.objects
        .filter(
            fondo=fondo,
            estado__in=(S.ESTADO_APROBADA, S.ESTADO_APROBADA_PARCIAL),
            valor_aprobado__gt=0,
            corte__isnull=True,
            fecha_respuesta__lt=fin,
            anonimizada_en__isnull=True,
        )
        .exclude(pedido__estado=Pedido.ESTADO_CANCELADO)
        .select_related("pedido__entrega", "finca")
        .order_by("finca__nombre", "nombre_completo", "id")
    )


def solicitudes_del_periodo(fondo, periodo):
    """(corte, solicitudes): si el mes ya se cerró devuelve lo congelado; si no, la vista previa."""
    corte = CorteNomina.objects.filter(fondo=fondo, periodo=periodo).first()
    if corte:
        return corte, corte.solicitudes.select_related("pedido__entrega", "finca").order_by("finca__nombre", "nombre_completo", "id")
    return None, solicitudes_pendientes_de_corte(fondo, periodo)


def resumen(solicitudes):
    datos = solicitudes.aggregate(total=Sum("valor_aprobado"))
    return Resumen(cantidad=solicitudes.count(), total=datos["total"] or Decimal("0"))


def _fecha_entrega(solicitud):
    entrega = getattr(solicitud.pedido, "entrega", None) if solicitud.pedido_id else None
    return timezone.localtime(entrega.entregado_en).strftime("%Y-%m-%d") if entrega else ""


def filas(solicitudes, periodo):
    for s in solicitudes:
        yield {
            "periodo": f"{periodo:%Y-%m}",
            "documento": s.documento,
            "nombre": s.nombre_completo,
            "empresa": s.empresa,
            "finca": str(s.finca) if s.finca_id else "",
            "numero_afiliado": s.numero_afiliado,
            "pedido": s.pedido.numero if s.pedido_id else "",
            "fecha_aprobacion": timezone.localtime(s.fecha_respuesta).strftime("%Y-%m-%d") if s.fecha_respuesta else "",
            "tipo": "Parcial" if s.estado == S.ESTADO_APROBADA_PARCIAL else "Total",
            "valor_pedido": s.valor_solicitado,
            "valor_descontar": s.valor_aprobado,
            "referencia_fondo": s.referencia_fondo,
            "fecha_entrega": _fecha_entrega(s),
            "solicitud": s.id,
        }


def _monto_csv(valor):
    # Excel en español (Colombia) usa coma decimal y punto y coma como separador.
    return f"{valor:.2f}".replace(".", ",")


def generar_csv(solicitudes, periodo):
    salida = io.StringIO()
    escritor = csv.writer(salida, delimiter=";")
    escritor.writerow([titulo for _, titulo in COLUMNAS])
    for fila in filas(solicitudes, periodo):
        escritor.writerow([
            _monto_csv(fila[c]) if isinstance(fila[c], Decimal) else fila[c]
            for c, _ in COLUMNAS
        ])
    return salida.getvalue().encode("utf-8-sig")


def generar_xlsx(solicitudes, periodo, fondo):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    libro = Workbook()
    hoja = libro.active
    hoja.title = f"Nómina {periodo:%Y-%m}"
    hoja.append([titulo for _, titulo in COLUMNAS])
    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="0B4D2C")

    total = Decimal("0")
    for fila in filas(solicitudes, periodo):
        hoja.append([fila[c] for c, _ in COLUMNAS])
        total += fila["valor_descontar"]

    hoja.append([])
    col_total = [c for c, _ in COLUMNAS].index("valor_descontar") + 1
    fila_total = [""] * len(COLUMNAS)
    fila_total[col_total - 2], fila_total[col_total - 1] = "TOTAL", total
    hoja.append(fila_total)
    hoja.cell(row=hoja.max_row, column=col_total - 1).font = Font(bold=True)
    hoja.cell(row=hoja.max_row, column=col_total).font = Font(bold=True)

    for columna in ("J", "K"):
        for celda in hoja[columna][1:]:
            celda.number_format = "#,##0.00"
    for indice, (_, titulo) in enumerate(COLUMNAS, start=1):
        hoja.column_dimensions[hoja.cell(row=1, column=indice).column_letter].width = max(14, len(titulo) + 4)
    hoja.freeze_panes = "A2"

    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


FORMATOS = {
    "csv": ("text/csv; charset=utf-8", "csv"),
    "xlsx": ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"),
}


def archivo_reporte(fondo, periodo, formato):
    """Devuelve (contenido, content_type, nombre_archivo)."""
    if formato not in FORMATOS:
        raise ReporteError("Formato no soportado. Use csv o xlsx.")
    _, solicitudes = solicitudes_del_periodo(fondo, periodo)
    contenido = (
        generar_xlsx(solicitudes, periodo, fondo) if formato == "xlsx" else generar_csv(solicitudes, periodo)
    )
    tipo, extension = FORMATOS[formato]
    nombre = f"nomina_{fondo.nit or fondo.id}_{periodo:%Y-%m}.{extension}"
    return contenido, tipo, nombre


def cerrar_corte(fondo, periodo, usuario=None):
    """Congela las solicitudes del mes en un corte. No se puede cerrar dos veces el mismo mes."""
    with transaction.atomic():
        if CorteNomina.objects.filter(fondo=fondo, periodo=periodo).exists():
            raise ReporteError(f"El corte de {periodo:%Y-%m} para {fondo} ya está cerrado.")
        ids = list(
            solicitudes_pendientes_de_corte(fondo, periodo)
            .select_for_update(of=("self",))
            .values_list("id", flat=True)
        )
        if not ids:
            raise ReporteError("No hay descuentos aprobados para cerrar en este periodo.")
        try:
            corte = CorteNomina.objects.create(fondo=fondo, periodo=periodo, generado_por=usuario)
        except IntegrityError:
            raise ReporteError(f"El corte de {periodo:%Y-%m} para {fondo} ya está cerrado.")
        S.objects.filter(id__in=ids).update(corte=corte)
        datos = resumen(corte.solicitudes.all())
        corte.cantidad, corte.total = datos.cantidad, datos.total
        corte.save(update_fields=["cantidad", "total"])
    return corte


def enviar_corte_por_correo(corte):
    """Envía el archivo del corte al correo del fondo y a sus gestores."""
    destinatarios = set(
        corte.fondo.gestores.filter(is_active=True).exclude(email="").values_list("email", flat=True)
    )
    if corte.fondo.correo:
        destinatarios.add(corte.fondo.correo)
    if not destinatarios:
        raise ReporteError("El fondo no tiene correo ni gestores con correo.")

    contenido, tipo, nombre = archivo_reporte(corte.fondo, corte.periodo, "xlsx")
    correo = EmailMessage(
        subject=f"Descuentos por nómina {corte.periodo:%Y-%m} · {corte.fondo}",
        body=(
            f"Adjuntamos el reporte de descuentos por nómina del periodo {corte.periodo:%Y-%m}.\n\n"
            f"Solicitudes: {corte.cantidad}\nTotal a descontar: ${corte.total:,.2f}\n\n"
            "El archivo contiene datos personales: trátelo conforme a la Ley 1581 de 2012."
        ),
        from_email=settings.DEFAULT_FROM_EMAIL,
        to=sorted(destinatarios),
    )
    correo.attach(nombre, contenido, tipo)
    correo.send(fail_silently=False)
    corte.enviado_en = timezone.now()
    corte.save(update_fields=["enviado_en"])
    return corte
